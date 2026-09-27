from io import BytesIO
from zipfile import ZipFile

from docx import Document

from app.evaluation_group_document import build_evaluation_group_docx
from app.evaluation_groups import EvaluationApplicant, EvaluationGroups


def applicant(identifier, name, number, grades, comments=None):
    return EvaluationApplicant(identifier, name, number, grades, comments or {})


def test_builds_ordered_grouped_applicant_docx_with_review_details_and_counts():
    groups = EvaluationGroups(
        buckets={
            "AAA": (applicant("1", "Alpha", "E-1", {"r1": "A", "r2": "A", "r3": "A"}),),
            "AAB": (),
            "ABB": (),
            "BBB": (),
            "BBC": (),
            "BCC": (),
            "CCC": (),
        },
        strong_discrepancy=(
            applicant("2", "Flagged", "E-2", {"r1": "A", "r2": "C"},
                      comments={"r1": "Clear match to the fellowship aims."}),
        ),
        awaiting_reviews=(
            applicant("2", "Flagged", "E-2", {"r1": "A", "r2": "C"}),
            applicant("3", "Pending", "E-3", {"r1": "B"}),
        ),
    )

    output = build_evaluation_group_docx(
        call_title="2026 Fellowship Call",
        call_code="CALL-2026",
        reviewer_roster=[("r1", "Ricky Reviewer"), ("r2", "Magda Reviewer"), ("r3", "Adriano Aguzzi")],
        grouping=groups,
    )

    assert isinstance(output, bytes)
    document = Document(BytesIO(output))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    assert text.index("AAA") < text.index("AAB") < text.index("ABB") < text.index("BBB")
    assert text.index("BBB") < text.index("BBC") < text.index("BCC") < text.index("CCC")
    assert "2026 Fellowship Call" in text
    assert "CALL-2026" in text
    assert "Ricky Reviewer (r1)" in text
    assert "Flagged" in text and "r1: A" in text and "r2: C" in text
    assert "Clear match to the fellowship aims." in text
    assert "2 awaiting reviews" in text
    assert "1 complete" in text
    assert "The partial A+C case appears in both awaiting reviews and strong discrepancies." in text


def test_docx_export_treats_untrusted_text_as_plain_text_without_hyperlinks():
    applicant_name = "<script>alert(1)</script> & [external](https://example.invalid)"
    groups = EvaluationGroups(
        buckets={key: () for key in ("AAA", "AAB", "ABB", "BBB", "BBC", "BCC", "CCC")},
        strong_discrepancy=(),
        awaiting_reviews=(applicant("x", applicant_name, "<number>", {"r1": "A"}),),
    )

    output = build_evaluation_group_docx(
        call_title="<Call>",
        call_code="<Code>",
        reviewer_roster={"r1": "<Reviewer 1>", "r2": "Reviewer 2", "r3": "Reviewer 3"},
        grouping=groups,
    )

    document = Document(BytesIO(output))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert applicant_name in text
    assert "<Call>" in text
    assert "<Code>" in text
    assert "<Reviewer 1> (r1)" in text
    with ZipFile(BytesIO(output)) as archive:
        assert b"w:hyperlink" not in archive.read("word/document.xml")
